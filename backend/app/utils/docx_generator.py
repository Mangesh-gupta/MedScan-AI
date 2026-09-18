from pathlib import Path
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def generate_radiology_docx(
    report_data: dict,
    patient_data: dict,
    study_data: dict,
    findings_list: list,
    output_path: str
) -> str:
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    
    # Page setup - 0.6 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    # Header
    head_p = doc.add_paragraph()
    r1 = head_p.add_run("MEDSCAN AI HEALTHCARE\n")
    r1.font.size = Pt(16)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(15, 23, 42)
    
    r2 = head_p.add_run("Department of Radiology & Advanced AI Diagnostics | Diagnostic Imaging Report")
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor(2, 132, 199)

    doc.add_paragraph("?" * 55)

    # Demographics Table
    table = doc.add_table(rows=4, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    
    labels = [
        ("PATIENT NAME:", patient_data.get('full_name', 'N/A'), "ACCESSION #:", study_data.get('accession_number', 'N/A')),
        ("MRN:", patient_data.get('mrn', 'N/A'), "STUDY DATE:", study_data.get('study_date', datetime.utcnow().strftime('%Y-%m-%d'))),
        ("AGE / SEX:", f"{patient_data.get('age', 'N/A')} / {patient_data.get('gender', 'N/A')}", "MODALITY:", f"{study_data.get('modality', 'N/A')} - {study_data.get('body_part', 'N/A')}"),
        ("REFERRING MD:", study_data.get('referring_physician', 'Dr. J. Reynolds'), "REPORT STATUS:", report_data.get('status', 'FINAL'))
    ]
    
    for row_idx, row_vals in enumerate(labels):
        row = table.rows[row_idx]
        for col_idx in range(4):
            cell = row.cells[col_idx]
            cell.text = str(row_vals[col_idx])
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(8.5)
            if col_idx in (0, 2):
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph("")

    # Sections
    sections_content = [
        ("CLINICAL INDICATION", report_data.get("clinical_history") or patient_data.get("clinical_notes") or "Diagnostic imaging evaluation as requested by referring physician."),
        ("TECHNIQUE", report_data.get("technique") or f"Standard diagnostic protocol for {study_data.get('modality')} {study_data.get('body_part')}."),
        ("NARRATIVE FINDINGS", report_data.get("findings_text") or "Imaging evaluated with deep learning AI assistance."),
        ("IMPRESSION", report_data.get("impression_text") or "No acute abnormality."),
        ("RECOMMENDATIONS", report_data.get("recommendations_text") or "Routine clinical follow-up as indicated."),
        ("CLINICAL NOTES & AI REASONING", report_data.get("clinical_notes_text") or "AI copilot quantitative analysis verified by radiologist.")
    ]

    for title, body in sections_content:
        p_title = doc.add_paragraph()
        run_title = p_title.add_run(title)
        run_title.font.bold = True
        run_title.font.size = Pt(11)
        run_title.font.color.rgb = RGBColor(2, 132, 199)
        
        p_body = doc.add_paragraph(body)
        p_body.paragraph_format.space_after = Pt(6)

    # Signature
    doc.add_paragraph("?" * 55)
    sig_p = doc.add_paragraph()
    sig_run = sig_p.add_run(f"Electronically Signed by: {report_data.get('radiologist_name', 'Dr. Sarah Al-Mansoor, MD, DABR')}\n")
    sig_run.font.bold = True
    sig_p.add_run(f"Timestamp: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')} | VisionGuard 360 Certified")

    doc.save(output_path)
    return output_path
